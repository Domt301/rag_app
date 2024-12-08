import unittest
import os
from file_handler import read_pdf, read_docx, read_txt, chunk_text, process_files

class TestFileHandler(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """
        Set up temporary test files for testing.
        """
        cls.test_dir = "test_files"
        os.makedirs(cls.test_dir, exist_ok=True)
        
        # Create test PDF file with actual text content
        cls.pdf_path = os.path.join(cls.test_dir, "test.pdf")
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="This is a test PDF file.", ln=True, align='L')
        pdf.output(cls.pdf_path)
        
        # Create test DOCX file
        cls.docx_path = os.path.join(cls.test_dir, "test.docx")
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is a test paragraph.")
        doc.save(cls.docx_path)
        
        # Create test TXT file
        cls.txt_path = os.path.join(cls.test_dir, "test.txt")
        with open(cls.txt_path, "w", encoding="utf-8") as f:
            f.write("This is a test text file.\nIt has multiple lines.\n")
    
    @classmethod
    def tearDownClass(cls):
        """
        Clean up temporary test files after testing.
        """
        for file_path in [cls.pdf_path, cls.docx_path, cls.txt_path]:
            if os.path.exists(file_path):
                os.remove(file_path)
        if os.path.exists(cls.test_dir):
            os.rmdir(cls.test_dir)

    def test_read_pdf(self):
        """
        Test reading a PDF file.
        """
        result = read_pdf(self.pdf_path)
        self.assertIn("This is a test PDF file.", result, "PDF content not read correctly.")
    
    def test_read_docx(self):
        """
        Test reading a DOCX file.
        """
        result = read_docx(self.docx_path)
        self.assertIn("This is a test paragraph.", result, "DOCX content not read correctly.")
    
    def test_read_txt(self):
        """
        Test reading a TXT file.
        """
        result = read_txt(self.txt_path)
        self.assertIn("This is a test text file.", result, "TXT content not read correctly.")
    
    def test_chunk_text(self):
        """
        Test chunking text.
        """
        text = "This is a test sentence. This is another sentence. Yet another sentence."
        chunks = chunk_text(text, max_length=20, chunk_overlap=5)
        self.assertGreater(len(chunks), 0, "Chunks were not created.")
        self.assertTrue(all(len(chunk) <= 20 for chunk in chunks), "Chunk size exceeds max_length.")

    def test_process_files(self):
        """
        Test processing files in a directory.
        """
        chunks = process_files(self.test_dir, chunk_size=50, chunk_overlap=10)
        self.assertGreater(len(chunks), 0, "No chunks were processed.")
        self.assertTrue(any("This is a test paragraph." in chunk for chunk in chunks),
                        "DOCX content not processed.")
        self.assertTrue(any("This is a test text file." in chunk for chunk in chunks),
                        "TXT content not processed.")

if __name__ == "__main__":
    unittest.main()